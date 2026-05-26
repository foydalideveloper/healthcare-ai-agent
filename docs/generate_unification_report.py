"""
Generate HRT+LRT Unified Database Migration Report as .docx
Complete technical report of what was done, how it works, and how teammates can use it.
Triple-H Co., Ltd. | April 15, 2026
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
    hs.font.name = 'Calibri'

doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(16)
doc.styles['Heading 3'].font.size = Pt(13)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)


def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Cm(1.27)


def code_block(code):
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'F5F5F5')
    shading.set(qn('w:val'), 'clear')
    p.paragraph_format.element.get_or_add_pPr().append(shading)


# ═══════════════════════════════════════════════════════════
# COVER PAGE
# ═══════════════════════════════════════════════════════════
for _ in range(5):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('HRT + LRT Unified Database\nMigration Complete Report')
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
run.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('What Was Done, How It Works, How To Use It')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x42, 0x42, 0x42)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Triple-H Co., Ltd. | April 15, 2026')
run.font.size = Pt(14)
run.bold = True
run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Project: LRT-HRT-Unified | ID: klnykuxzucujahucvbct\nRegion: ap-northeast-2 (Seoul) | Status: ACTIVE_HEALTHY')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 1. WHAT WAS DONE
# ═══════════════════════════════════════════════════════════
doc.add_heading('1. What Was Done', level=1)

doc.add_paragraph(
    'We created a NEW unified Supabase project that merges the Healthcare (HRT) database '
    '(14 tables, 5,500+ rows of real data) and the Lifetime (LRT) database (17 tables, schema only) '
    'into ONE database with 28 unified tables. Both old projects remain untouched as backups.'
)

doc.add_heading('1.1 Migration Steps Completed', level=2)

add_table(
    ['Step', 'What Was Done', 'Result'],
    [
        ['1', 'Created new Supabase project "LRT-HRT-Unified"', 'Project ID: klnykuxzucujahucvbct, Seoul region, $10/month'],
        ['2', 'Created 28 unified tables with BIGINT PKs', 'All conflicts resolved: no UUIDs as PK, no ENUMs, gender=X'],
        ['3', 'Created 33 indexes + pgvector HNSW', 'Composite time-range, GIN JSONB, partial, vector similarity'],
        ['4', 'Migrated standard health data (424 rows)', '26 categories + 298 norms + 44 weights + 56 lifestyle plans'],
        ['5', 'Migrated 8 users with UUID\u2192BIGINT mapping', 'Old UUIDs preserved as user_token for API compatibility'],
        ['6', 'Migrated 5,042 biometric records', 'Heart rate, glucose, SpO2 data \u2014 all FKs updated to BIGINT'],
        ['7', 'Migrated 193 activity events + 1 lifestyle + 1 health level', 'All glasses pipeline test data preserved'],
        ['8', 'Created 45 RLS policies', 'Per-user data isolation for HIPAA/PIPA compliance'],
        ['9', 'Updated FastAPI backend (.env, models, endpoints)', 'Backend now points to unified project'],
        ['10', 'Tested all API endpoints', 'Users, events, food lookup \u2014 all working correctly'],
    ])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 2. CONNECTION DETAILS
# ═══════════════════════════════════════════════════════════
doc.add_heading('2. Connection Details (For All Team Members)', level=1)

doc.add_heading('2.1 New Unified Project', level=2)

add_table(
    ['Property', 'Value'],
    [
        ['Project Name', 'LRT-HRT-Unified'],
        ['Project ID', 'klnykuxzucujahucvbct'],
        ['URL', 'https://klnykuxzucujahucvbct.supabase.co'],
        ['Region', 'ap-northeast-2 (Seoul)'],
        ['PostgreSQL Version', '17.6'],
        ['Anon Key', 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImtsbnlrdXh6dWN1amFodWN2YmN0Iiwicm9sZSI6ImFub24i...'],
        ['Dashboard', 'https://supabase.com/dashboard/project/klnykuxzucujahucvbct'],
    ])

doc.add_paragraph(
    'Important: Get the full anon key and service role key from the Supabase dashboard: '
    'Settings \u2192 API \u2192 Project API keys. Never commit the service role key to Git.'
)

doc.add_heading('2.2 Old Projects (Backups \u2014 Do Not Delete)', level=2)

add_table(
    ['Project', 'ID', 'Status', 'Data'],
    [
        ['healthcare-ai-agent (HRT)', 'oqotdxlmdgjieukyzegj', 'Keep as backup', '14 tables, 5,500+ rows'],
        ['LRT', 'rietbmwbraqisoneoill', 'Keep as backup', '17 tables, 0 rows (schema only)'],
    ])

doc.add_heading('2.3 How To Connect Your Backend', level=2)

doc.add_paragraph('Add these to your .env file:')

code_block(
    '# UNIFIED PROJECT\n'
    'SUPABASE_URL=https://klnykuxzucujahucvbct.supabase.co\n'
    'SUPABASE_ANON_KEY=<get from dashboard>\n'
    'SUPABASE_SERVICE_ROLE_KEY=<get from dashboard>'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 3. DATABASE SCHEMA
# ═══════════════════════════════════════════════════════════
doc.add_heading('3. Database Schema (28 Tables)', level=1)

doc.add_heading('3.1 Table Overview by Domain', level=2)

add_table(
    ['#', 'Table', 'Domain', 'Rows', 'Purpose'],
    [
        ['', '', 'SHARED CORE', '', ''],
        ['1', 'users', 'Shared', '8', 'User profiles (BIGINT PK + UUID token)'],
        ['2', 'demographic_clusters', 'Shared', '0', 'Population segments for clustering'],
        ['3', 'agents', 'Shared', '0', 'AI agents (TEXT role, not ENUM)'],
        ['4', 'agent_sessions', 'Shared', '0', 'Agent conversation sessions'],
        ['5', 'multimodal_inputs', 'Shared', '0', 'Photos, audio, video with pgvector + JSONB'],
        ['', '', 'HEALTH DOMAIN', '', ''],
        ['6', 'std_population_category', 'Health', '26', 'Age/gender/disability categories'],
        ['7', 'std_diagnosis_norm', 'Health', '298', 'Normal/warning/critical health ranges'],
        ['8', 'std_lifestyle_plan', 'Health', '56', 'Diet/exercise/sleep/medication plans (JSONB)'],
        ['9', 'std_disease_risk_weight', 'Health', '44', 'Disease risk contribution weights'],
        ['10', 'user_biometric', 'Health', '5,042', 'Heart rate, glucose, SpO2 from wearables'],
        ['11', 'user_diagnosis', 'Health', '0', 'Medical lab data (BMI, BP, cholesterol)'],
        ['12', 'user_lifestyle', 'Health', '1', 'Daily diet, exercise, sleep aggregation'],
        ['13', 'user_health_level', 'Health', '1', 'Healthcare Index (HCI), Disease Risk Index'],
        ['14', 'simulation_result', 'Health', '0', 'LSTM-Transformer health predictions'],
        ['', '', 'LIFE DOMAIN', '', ''],
        ['15', 'standard_lrt_header', 'Life', '0', 'Standard LRT templates per cluster'],
        ['16', 'standard_lrt_events', 'Life', '0', 'Expected life events (marriage, retirement)'],
        ['17', 'standard_lrt_activities', 'Life', '0', 'Expected behaviors per event'],
        ['18', 'standard_lrt_products', 'Life', '0', 'Product recommendations per activity'],
        ['19', 'personal_lrt_header', 'Life', '0', 'Per-user personalized LRT'],
        ['20', 'personal_lrt_events', 'Life', '0', 'Individual predicted life events'],
        ['21', 'personal_lrt_activities', 'Life', '0', 'Individual behaviors with spend'],
        ['22', 'lrt_simulations', 'Life', '0', 'Monte Carlo life event simulations'],
        ['', '', 'CROSS-DOMAIN', '', ''],
        ['23', 'user_activity_event', 'Cross', '193', 'AI glasses events (food + purchases + exercise)'],
        ['24', 'agent_relationship_permissions', 'Cross', '0', '1st/2nd/3rd degree data sharing'],
        ['25', 'agent_conversation_log', 'Cross', '0', 'All agent chat history'],
        ['26', 'history_record_table', 'Cross', '0', 'Blockchain-style audit trail'],
        ['27', 'contracts', 'Cross', '0', 'Purchase contracts (LRT pre-order)'],
        ['28', 'orders', 'Cross', '0', 'Order execution with delivery'],
    ])

doc.add_page_break()

doc.add_heading('3.2 Users Table (Unified)', level=2)

doc.add_paragraph(
    'This is the most important table. Every other table references it via user_id (BIGINT). '
    'The old UUID is preserved as user_token for external API use.'
)

add_table(
    ['Column', 'Type', 'Purpose', 'Source'],
    [
        ['user_id', 'BIGINT (PK, auto)', 'Internal ID for all JOINs and FKs', 'New'],
        ['user_token', 'UUID (unique)', 'External ID for APIs, mobile app, glasses', 'New'],
        ['user_hash', 'VARCHAR(64)', 'SHA-256 de-identification key', 'HRT'],
        ['auth_user_id', 'UUID', 'Link to Supabase Auth (auth.users)', 'LRT'],
        ['cluster_id', 'BIGINT FK', 'Demographic cluster for LRT predictions', 'LRT'],
        ['birth_year', 'SMALLINT', 'Exact birth year (more precise than age_group)', 'LRT'],
        ['age_group', 'VARCHAR(20)', 'Age group (backward compatibility)', 'HRT'],
        ['gender', "CHAR(1) CHECK (M,F,X)", 'X = ICAO standard for non-binary', 'Unified'],
        ['nationality', 'CHAR(3)', 'ISO country code (was country_code in HRT)', 'Both'],
        ['disability_yn', 'BOOLEAN', 'Has disability?', 'HRT'],
        ['disability_type', 'VARCHAR(50)', 'Type of disability (was in user_demographic)', 'HRT'],
        ['income_decile', 'SMALLINT (1-10)', 'Income decile', 'HRT'],
        ['past_history', 'TEXT[]', 'Past medical history array (was in user_demographic)', 'HRT'],
        ['family_history', 'TEXT[]', 'Family medical history array (was in user_demographic)', 'HRT'],
        ['occupation_category', 'TEXT', 'Job category', 'LRT'],
        ['household_type', 'TEXT', 'Household composition', 'LRT'],
        ['income_bracket', 'TEXT', 'Income bracket text', 'LRT'],
        ['region', 'TEXT', 'Geographic region', 'LRT'],
        ['mbti', 'CHAR(4)', 'MBTI personality type', 'LRT'],
        ['is_active', 'BOOLEAN', 'Account active?', 'HRT'],
        ['created_at', 'TIMESTAMPTZ', 'Registration time', 'Both'],
        ['deleted_at', 'TIMESTAMPTZ', 'Soft delete time', 'LRT'],
    ])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 4. KEY DECISIONS
# ═══════════════════════════════════════════════════════════
doc.add_heading('4. Key Decisions Made During Unification', level=1)

add_table(
    ['Decision', 'Old HRT', 'Old LRT', 'Unified', 'Why'],
    [
        ['Primary Key', 'UUID', 'BIGINT', 'BIGINT + UUID token', 'BIGINT is faster for JOINs. UUID token used externally for security.'],
        ['Gender Code', "CHECK (M,F,O)", "ENUM (M,F,X)", "CHECK (M,F,X)", "X is ICAO standard. CHECK is more flexible than ENUM."],
        ['Agent Roles', 'None (simple log)', "ENUM (4 roles)", "TEXT (unlimited)", "11+ agent roles needed. TEXT avoids ALTER TYPE for new roles."],
        ['Multimodal', 'pgvector + VARCHAR', 'JSONB + TTL', 'pgvector + JSONB + TTL + context', 'Best features from both combined.'],
        ['Input Types', "CHECK (voice/image/video/sensor)", "ENUM (5 types)", "CHECK (6 types incl. VIDEO)", "Added VIDEO type that LRT was missing."],
        ['Simulation', "simulation_result", "lrt_simulations", "Both kept (no rename)", "Different purposes. Renaming breaks HRT code."],
        ['user_demographic', 'Separate table', 'N/A', 'Merged into users', "Fields (past_history, family_history) added to users table."],
        ['Processing status', "CHECK (3 values)", "N/A", "CHECK (4 values)", "Added 'needs_cloud_review' status."],
    ])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 5. UUID TO BIGINT MAPPING
# ═══════════════════════════════════════════════════════════
doc.add_heading('5. UUID to BIGINT Mapping', level=1)

doc.add_paragraph(
    'All 8 existing HRT users were migrated. Their old UUID is now the user_token column. '
    'All 9 child tables (biometric, diagnosis, lifestyle, events, health_level, multimodal, '
    'conversation_log, simulation_result, activity_event) now use BIGINT user_id.'
)

add_table(
    ['Old UUID (user_token)', 'New BIGINT (user_id)', 'Gender', 'Age Group'],
    [
        ['5ed9da6d-6b9c-4a5f-bd80-587f0e54280d', '1', 'M', '30-39'],
        ['7fc5ab71-0e56-419c-933b-b810fa0e2d65', '2', 'M', '30-39'],
        ['35ce8c78-1d89-41d1-b2c8-6245861ba69c', '3', 'M', '30-39'],
        ['508f0c7d-b7e3-4a37-9d6e-6117446d5de5', '4', 'M', '50-59'],
        ['0aaac176-d427-49f8-8e10-79f4053820f0', '5', 'F', '40-49'],
        ['d294cbfd-4247-48d3-b619-805e5eb370ef', '6', 'M', '70-79'],
        ['d93f2941-b202-479a-82e7-20b08a69395d', '7', 'F', '19-29'],
        ['2b513e01-8f3a-4315-ac00-4cd0ac342bed', '8', 'M', '30-39'],
    ])

doc.add_paragraph(
    'The _migration_uuid_map table exists in the database for reference. It can be dropped '
    'after all team members confirm their code is updated.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 6. RLS POLICIES
# ═══════════════════════════════════════════════════════════
doc.add_heading('6. Row Level Security (45 Policies)', level=1)

doc.add_paragraph(
    'Every table has RLS enabled. Users can only see and modify their own data. '
    'Standard reference tables are read-only for authenticated users.'
)

doc.add_heading('6.1 How RLS Works', level=2)

doc.add_paragraph(
    'A helper function get_my_user_id() converts the Supabase Auth UUID (auth.uid()) '
    'to the BIGINT user_id used in all tables. Every policy calls this function.'
)

code_block(
    '-- This function is already created in the database\n'
    'CREATE FUNCTION get_my_user_id() RETURNS BIGINT AS $$\n'
    '    SELECT user_id FROM public.users WHERE auth_user_id = auth.uid()\n'
    '$$ LANGUAGE sql SECURITY DEFINER STABLE;'
)

doc.add_heading('6.2 Policy Summary', level=2)

add_table(
    ['Table Type', 'SELECT', 'INSERT', 'UPDATE', 'Example'],
    [
        ['users', 'Own profile', '\u2014', 'Own profile', 'Users can read and update only their own profile'],
        ['Standard tables (9)', 'Authenticated', '\u2014', '\u2014', 'Anyone logged in can read health norms, disease weights, etc.'],
        ['Health personal (8)', 'Own data', 'Own data', 'activity_event only', 'User 5 cannot see User 3\'s biometric data'],
        ['LRT personal (4)', 'Own data', 'Own data', '\u2014', 'User\'s own life predictions and activities'],
        ['Cross-domain (6)', 'Own agents/contracts', 'Own', '\u2014', 'Permissions, conversation logs, audit trail'],
    ])

doc.add_paragraph(
    'Important for LRT team: When creating new records, the user_id must match '
    'get_my_user_id(). If using the service role key (bypasses RLS), you must still '
    'set user_id correctly.'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 7. HOW TO USE (FOR LRT TEAM)
# ═══════════════════════════════════════════════════════════
doc.add_heading('7. How To Use This Database (For LRT Team)', level=1)

doc.add_heading('7.1 Your Tables Are Ready', level=2)

doc.add_paragraph(
    'All 17 LRT tables from your old project have been recreated in the unified database '
    'with these changes:'
)

bullet('user_id is now BIGINT (was BIGINT in your old project too \u2014 no change needed)')
bullet('Gender uses CHAR(1) CHECK instead of ENUM \u2014 update any code that references the ENUM type')
bullet('Agent roles use TEXT instead of ENUM \u2014 update code accordingly')
bullet('All other ENUM types (contract_status, payment_method, etc.) are now TEXT with CHECK constraints')
bullet('multimodal_inputs merges your old table with HRT\'s \u2014 now includes pgvector embedding_vec and context fields')

doc.add_heading('7.2 What You Need To Change In Your Code', level=2)

add_table(
    ['Change', 'Old LRT Code', 'New Unified Code'],
    [
        ['Supabase URL', 'https://rietbmwbraqisoneoill.supabase.co', 'https://klnykuxzucujahucvbct.supabase.co'],
        ['Gender type', "ENUM gender_code ('M','F','X')", "CHAR(1) CHECK ('M','F','X')"],
        ['Agent role', "ENUM agent_role_code", "TEXT (any value)"],
        ['Contract type', "ENUM contract_type_code", "TEXT CHECK ('SHORT','LONG','SUBSCRIPTION')"],
        ['Simulation trigger', "ENUM simulation_trigger_type_code", "TEXT CHECK ('ANOMALY','PEER_UPDATE','GROUP_SHIFT','MANUAL')"],
        ['All other ENUMs', 'PostgreSQL ENUM types', 'TEXT with CHECK constraints (same values)'],
        ['multimodal table', 'multimodal_inputs (no pgvector)', 'multimodal_inputs (has embedding_vec vector(1536))'],
    ])

doc.add_heading('7.3 Creating New Users', level=2)

doc.add_paragraph('When your LRT backend creates a new user:')

code_block(
    'INSERT INTO users (\n'
    '    birth_year, gender, nationality, occupation_category,\n'
    '    household_type, income_bracket, region, mbti, cluster_id\n'
    ') VALUES (\n'
    '    1990, \'M\', \'KOR\', \'IT\',\n'
    '    \'single\', \'middle\', \'Seoul\', \'INTJ\', 1\n'
    ') RETURNING user_id, user_token;\n'
    '\n'
    '-- Returns: user_id=9, user_token=<auto-generated UUID>\n'
    '-- Use user_id (BIGINT) for all internal operations\n'
    '-- Use user_token (UUID) for APIs and external communication'
)

doc.add_heading('7.4 Connecting Health and Life Data', level=2)

doc.add_paragraph(
    'The power of unification: you can now query health data alongside life data '
    'for the same user.'
)

code_block(
    '-- Example: Get user\'s health status + predicted life events\n'
    'SELECT u.user_id, u.age_group, u.mbti,\n'
    '       hl.healthcare_index, hl.disease_risk_json,\n'
    '       pe.event_name, pe.predicted_date, pe.confidence_score\n'
    'FROM users u\n'
    'LEFT JOIN user_health_level hl ON u.user_id = hl.user_id\n'
    'LEFT JOIN personal_lrt_header ph ON u.user_id = ph.user_id\n'
    'LEFT JOIN personal_lrt_events pe ON ph.plrt_id = pe.plrt_id\n'
    'WHERE u.user_id = 1;'
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 8. DATA VERIFICATION
# ═══════════════════════════════════════════════════════════
doc.add_heading('8. Data Verification', level=1)

doc.add_paragraph('All data was verified after migration:')

add_table(
    ['Data', 'Old Project Count', 'New Unified Count', 'Match'],
    [
        ['Population categories', '26', '26', 'YES \u2713'],
        ['Diagnosis norms', '298', '298', 'YES \u2713'],
        ['Disease risk weights', '44', '44', 'YES \u2713'],
        ['Lifestyle plans', '56', '56', 'YES \u2713'],
        ['Users', '8', '8', 'YES \u2713'],
        ['Biometric records', '5,042', '5,042', 'YES \u2713'],
        ['Activity events', '193', '193', 'YES \u2713'],
        ['Lifestyle records', '1', '1', 'YES \u2713'],
        ['Health level records', '1', '1', 'YES \u2713'],
        ['Total data rows', '5,669', '5,669', 'YES \u2713 (zero data lost)'],
    ])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 9. INDEXES
# ═══════════════════════════════════════════════════════════
doc.add_heading('9. Indexes (33 Total)', level=1)

add_table(
    ['Index Name', 'Table', 'Columns', 'Type'],
    [
        ['idx_biometric_user_time', 'user_biometric', 'user_id, measured_at DESC', 'Composite'],
        ['idx_diagnosis_user_time', 'user_diagnosis', 'user_id, period_start DESC', 'Composite'],
        ['idx_lifestyle_user_time', 'user_lifestyle', 'user_id, recorded_date DESC', 'Composite'],
        ['idx_health_level_user_time', 'user_health_level', 'user_id, measured_at DESC', 'Composite'],
        ['idx_multimodal_user_time', 'multimodal_inputs', 'user_id, collected_at DESC', 'Composite'],
        ['idx_conversation_user_time', 'agent_conversation_log', 'user_id, created_at DESC', 'Composite'],
        ['idx_simulation_user_time', 'simulation_result', 'user_id, run_at DESC', 'Composite'],
        ['idx_activity_user_type_time', 'user_activity_event', 'user_id, event_type, detected_at DESC', 'Composite'],
        ['idx_activity_unverified', 'user_activity_event', 'WHERE verified=false AND confidence<0.7', 'Partial'],
        ['idx_lifestyle_medication_gin', 'user_lifestyle', 'medication_json', 'GIN jsonb_path_ops'],
        ['idx_disease_risk_gin', 'user_health_level', 'disease_risk_json', 'GIN jsonb_path_ops'],
        ['idx_activity_data_gin', 'user_activity_event', 'structured_data', 'GIN jsonb_path_ops'],
        ['idx_multimodal_embedding', 'multimodal_inputs', 'embedding_vec', 'HNSW (m=24, ef=128)'],
        ['+ 20 more', 'Various', 'Device, period, scenario, agent, LRT tables', 'Various'],
    ])

doc.add_page_break()

# ═══════════════════════════════════════════════════════════
# 10. WHAT'S NEXT
# ═══════════════════════════════════════════════════════════
doc.add_heading('10. What\'s Next', level=1)

add_table(
    ['Task', 'Owner', 'Priority', 'Notes'],
    [
        ['LRT team: point backend to new project', 'LRT team', 'HIGH', 'Change Supabase URL + keys in .env'],
        ['LRT team: update ENUM references to TEXT', 'LRT team', 'HIGH', 'gender_code, agent_role_code, etc.'],
        ['LRT team: populate standard_lrt_* tables', 'LRT team', 'HIGH', 'These tables are empty \u2014 need data'],
        ['LRT team: populate demographic_clusters', 'LRT team', 'MEDIUM', 'Needed for user clustering'],
        ['HRT team: test all existing endpoints', 'HRT team', 'DONE', 'Tested and working'],
        ['Both: link Supabase Auth (auth_user_id)', 'Both', 'MEDIUM', 'Existing 8 users need auth_user_id backfill when they register'],
        ['Both: set up Supabase Storage buckets', 'Both', 'MEDIUM', 'For thumbnails and media clips'],
        ['Both: clean up _migration_uuid_map table', 'Both', 'LOW', 'Can be dropped after code is updated'],
        ['Both: start cross-domain features', 'Both', 'FUTURE', 'Health data informing LRT predictions and vice versa'],
    ])

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Triple-H Co., Ltd. | HRT+LRT Unified Database Migration Report | April 15, 2026\n'
                'Project: klnykuxzucujahucvbct | Region: Seoul | Status: ACTIVE_HEALTHY')
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

# Save
output_path = r'C:\Users\tripleh\projects\healthcare-ai-agent\docs\HRT_LRT_Migration_Complete_Report.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
